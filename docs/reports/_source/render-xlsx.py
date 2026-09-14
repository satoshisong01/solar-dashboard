"""FIRST C&D company-style workbook from a JSON spec.

Usage: py -3.11 render-xlsx.py book.json out.xlsx
sheet: {"name", "title", "description", "columns": [{"header", "width", "align"}], "rows": [[...]],
        "freeze": true, "filter": true, "notes": [...]}
Cell values starting with "=" are written as formulas; http(s) URLs become hyperlinks.
"""
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

FONT = "Pretendard"
INK, BLUE, MUTED, HEAD, RULE, TINT = "222222", "4081ED", "858585", "DCE6F2", "D9DEE6", "EEF3FC"
thin = Side(style="thin", color=RULE)
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)


def write_sheet(ws, sh, meta):
    ws.sheet_view.showGridLines = False
    ncol = len(sh["columns"])
    last = get_column_letter(ncol)
    ws["A1"] = sh["title"]
    ws["A1"].font = Font(name=FONT, size=16, bold=True, color=BLUE)
    ws.row_dimensions[1].height = 30
    ws["A2"] = sh.get("description", "")
    ws["A2"].font = Font(name=FONT, size=10, color=MUTED)
    ws.merge_cells(f"A2:{last}2")
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[2].height = sh.get("descriptionHeight", 34)
    ws["A3"] = f"{meta['company']} · 작성일 {meta['date']}"
    ws["A3"].font = Font(name=FONT, size=9, color=MUTED)

    header_row = 5
    for i, col in enumerate(sh["columns"], start=1):
        c = ws.cell(row=header_row, column=i, value=col["header"])
        c.font = Font(name=FONT, size=10, bold=True, color=INK)
        c.fill = PatternFill("solid", start_color=HEAD)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
        ws.column_dimensions[get_column_letter(i)].width = col.get("width", 16)
    ws.row_dimensions[header_row].height = 26

    for r, row in enumerate(sh["rows"], start=header_row + 1):
        group = isinstance(row, dict)
        values = row["cells"] if group else row
        for i, v in enumerate(values, start=1):
            col = sh["columns"][i - 1]
            c = ws.cell(row=r, column=i, value=v)
            c.font = Font(name=FONT, size=10, color=INK, bold=group)
            if group:
                c.fill = PatternFill("solid", start_color=TINT)
            c.alignment = Alignment(horizontal=col.get("align", "left"), vertical="center", wrap_text=True)
            c.border = BORDER
            if isinstance(v, str) and v.startswith(("http://", "https://")):
                c.hyperlink = v
                c.font = Font(name=FONT, size=9, color=BLUE, underline="single")
            if isinstance(v, str) and v in ("●", "◐"):
                c.font = Font(name=FONT, size=11, bold=True, color=BLUE if v == "●" else MUTED)
        if "rowHeight" in sh:
            ws.row_dimensions[r].height = sh["rowHeight"]

    end = header_row + len(sh["rows"])
    if sh.get("freeze", True):
        ws.freeze_panes = ws.cell(row=header_row + 1, column=sh.get("freezeCols", 1) + 1)
    if sh.get("filter", True) and sh["rows"]:
        ws.auto_filter.ref = f"A{header_row}:{last}{end}"
    note_row = end + 2
    for n in sh.get("notes", []):
        ws.cell(row=note_row, column=1, value=n).font = Font(name=FONT, size=9, color=MUTED)
        note_row += 1
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToHeight = 0


def main():
    spec = json.loads(Path(sys.argv[1]).read_text(encoding="utf-8"))
    wb = Workbook()
    wb.remove(wb.active)
    for sh in spec["sheets"]:
        write_sheet(wb.create_sheet(sh["name"]), sh, spec["meta"])
    wb.properties.creator = spec["meta"]["company"]
    wb.properties.title = spec["meta"]["title"]
    wb.save(sys.argv[2])
    print("wrote", sys.argv[2], len(spec["sheets"]), "sheets")


if __name__ == "__main__":
    main()
